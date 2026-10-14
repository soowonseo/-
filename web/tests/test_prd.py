import sys, base64, io, json, csv
sys.path.insert(0, __import__("os").path.dirname(__file__))
from common import *
from pathlib import Path
from playwright.sync_api import sync_playwright
from docx import Document
import openpyxl
A = OUT / "artifact"
page_html = '<!doctype html><html><head><meta charset=utf8><meta name=viewport content="width=device-width,initial-scale=1,viewport-fit=cover"></head><body>' + (A / "manual-writer.html").read_text(encoding="utf-8") + "</body></html>"
stub = """
window.__saved = null;
window.claude = { use: async (n) => {
  if (n === 'downloads') return { save: async ({filename, data}) => { const b = new Uint8Array(await data.arrayBuffer()); let s=''; for (let i=0;i<b.length;i+=0x8000) s += String.fromCharCode.apply(null, b.subarray(i,i+0x8000)); window.__saved = {filename, b64: btoa(s)}; return {status:'saved'}; } };
  return null; } };
"""
OUT = A / "out"; OUT.mkdir(exist_ok=True)

def saved(pg, name):
    pg.evaluate("window.__saved = null"); pg.click("#download"); pg.wait_for_function("window.__saved", timeout=30000)
    s = pg.evaluate("window.__saved"); data = base64.b64decode(s["b64"]); (OUT / name).write_bytes(data); return s["filename"], data

def metrics(pg): return pg.locator(".metrics").inner_text().replace("\n", " ")
def confirm(pg, author="담당자A", reviewer="파트장C", summary=None):
    pg.fill("#author", author); pg.fill("#reviewer", reviewer)
    if summary: pg.fill("#summary", summary)
    pg.click("#confirm"); pg.wait_for_selector("#download", timeout=60000)
def preset(pg, label):
    pg.click("[data-page=input]"); use_preset(pg, label); pg.wait_for_timeout(800)
def generate(pg):
    pg.click("#generate")
    try: pg.wait_for_selector(".compare", timeout=20000)
    except Exception:
        print("!! 초안 실패:", pg.locator(".notice.warn").all_inner_texts(), "| 버튼:", pg.locator("#generate").is_disabled() if pg.locator("#generate").count() else "없음", "| busy:", pg.locator("#stream").all_inner_texts(), "| 문제:", pg.locator(".problems").all_inner_texts()); pg.screenshot(path=str(A / "fail.png"), full_page=True); raise
def restart(pg): pg.click("#restart"); pg.wait_for_selector("#generate")
def snap(ws):
    return ({c.coordinate: (c.value, bool(c.font.b)) for row in ws.iter_rows() for c in row if c.value is not None},
            [(i.anchor._from.row, round(i.anchor.ext.width / 9525), round(i.anchor.ext.height / 9525)) for i in ws._images],
            {r: ws.row_dimensions[r].height for r in range(1, ws.max_row + 1)})
def heads(doc): return [f"{p.style.name[-1]}:{p.text}" for p in doc.paragraphs if p.style.name.startswith("Heading")]

with sync_playwright() as p:
    b = launch(p)
    pg = b.new_page(viewport={"width": 1300, "height": 900})
    errors = []; pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: m.type == "error" and print("console:", m.text))
    route_libs(pg)
    pg.route("http://local.test/", lambda r: r.fulfill(body=page_html, content_type="text/html"))
    pg.add_init_script(stub)
    pg.goto("http://local.test/"); pg.wait_for_selector("#generate"); pg.wait_for_timeout(800)
    print("첫 화면(빈 화면): 파일", pg.locator(".file-row, .img-card").count(), "| 안내", pg.locator(".start-hint").count(), "| 버튼", "비활성" if pg.locator("#generate").is_disabled() else "활성")
    use_preset(pg, "fund"); pg.wait_for_timeout(500)
    print("예시 채운 뒤: 파일", pg.locator(".file-row, .img-card").count(), "| 형식", pg.evaluate("state.format"))

    # T7 신규 엑셀 (펀드변경 원자료)
    pg.check("input[name=format][value=xlsx]"); generate(pg); confirm(pg)
    name, data = saved(pg, "t7.xlsx"); ws = openpyxl.load_workbook(io.BytesIO(data)).active
    e = openpyxl.load_workbook(R + "mock/예상결과_펀드변경_매뉴얼_v2_단독.xlsx").active
    a, x = snap(ws), snap(e)
    diff = sorted(k for k in set(a[0]) | set(x[0]) if a[0].get(k) != x[0].get(k))
    print("[T7 신규 엑셀]", name, "| 시트", ws.title, "| 목표와 다른 칸", diff, "| 그림 같음", a[1] == x[1], "| 행높이 15", set(a[2].values()))
    restart(pg)

    # T1 신규 Word + 4-3 피드백
    preset(pg, "신규: 펀드변경 원자료"); generate(pg)
    pg.locator("#fb-6 summary").click(); pg.fill("#fbc-6", "오류 코드별 조치 추가")
    t = pg.input_value("#fbt-6"); pg.fill("#fbt-6", t.replace("메시지에 따라 수정한다.", "메시지에 따라 수정한다. (E-102 합계, E-105 판매중지)")); pg.click("[data-apply='6']")
    confirm(pg); name, data = saved(pg, "t1.docx"); d = Document(io.BytesIO(data))
    print("[T1 신규 Word]", name, "| 그림", len(d.inline_shapes), "| 개정", [c.text for c in d.tables[0].rows[1].cells], "| 4-3 반영", any("E-105 판매중지" in q.text for q in d.paragraphs))
    restart(pg)

    # T2 줄글만 → 캡처 1곳 업로드 → 경고 후 확정
    preset(pg, "신규: 자동재배분 러프 메모"); generate(pg)
    print("[T2 줄글만]", metrics(pg))
    pg.set_input_files("#cap-4-0", R + "원자료/캡처1_신청화면.png"); pg.wait_for_timeout(500)
    print("   업로드 후", metrics(pg), "| 경고", pg.locator(".notice.warn").first.inner_text()[:40])
    confirm(pg); name, data = saved(pg, "t2.docx"); d = Document(io.BytesIO(data))
    txt = [q.text for q in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells]
    print("   Word: 그림", len(d.inline_shapes), "| [캡처 필요]", sum(x.startswith("[캡처 필요") for x in txt), "| [확인 필요]", sum(x.count("[확인 필요]") for x in txt))
    restart(pg)

    # T3 Word 갱신
    preset(pg, "갱신: 기존 Word v1 + 규정 변경 메모"); generate(pg)
    cols = pg.locator(".compare > .stack"); asis, tobe = cols.nth(0).inner_text(), cols.nth(1).inner_text()
    print("[T3 Word 갱신]", metrics(pg), "| As-Is 6회·1,000원·다음 영업일", all(k in asis for k in ["6회", "1,000원", "다음 영업일"]), "| To-Be 12회·2,000원·+2영업일·판매중지", all(k in tobe for k in ["12회", "2,000원", "+2영업일", "판매중지"]), "| As-Is 그림", cols.nth(0).locator("img").count())
    confirm(pg, summary="규정 현행화"); name, data = saved(pg, "t3.docx"); d = Document(io.BytesIO(data))
    print("   ", name, "| 개정이력", [[c.text[:12] for c in r.cells][:3] for r in d.tables[0].rows[1:]])
    restart(pg)

    # T4 엑셀 갱신
    preset(pg, "갱신: 기존 엑셀 v1 '펀드변경' 시트"); generate(pg)
    asis = pg.locator(".compare > .stack").nth(0)
    print("[T4 엑셀 갱신]", metrics(pg), "| As-Is 그림", asis.locator("img").count(), "| 옆 설명", "현재비율 옆 칸에 변경비율 입력" in asis.inner_text(), "| 아래 설명", "기록용 엑셀. 판정 칸 확인" in asis.inner_text(), "| 도형 안내", "도형" in asis.inner_text())
    confirm(pg, summary="규정 현행화, 판매중지·적합성 확인 절차 추가, 캡처·쿼리 보강"); name, data = saved(pg, "t4.xlsx")
    o = openpyxl.load_workbook(io.BytesIO(data)); s = openpyxl.load_workbook(R + "excel/가상_업무매뉴얼_v1.xlsx"); e = openpyxl.load_workbook(R + "mock/예상결과_업무매뉴얼_v2.xlsx")
    print("   ", name, "| 시트", o.sheetnames)
    for n in ["목차", "주소변경"]: print(f"    '{n}' 원본과 같음:", snap(o[n]) == snap(s[n]))
    a, x = snap(o["펀드변경"]), snap(e["펀드변경"])
    print("    '펀드변경' 목표와 다른 칸:", sorted(k for k in set(a[0]) | set(x[0]) if a[0].get(k) != x[0].get(k)), "| 그림 같음", a[1] == x[1])
    restart(pg)

    # T5 Word 업무 추가
    preset(pg, "추가: Word 1~6장 + 변액 자동재배분"); generate(pg)
    print("[T5 Word 추가]", metrics(pg), "| As-Is 목차에 7장:", "7. 변액 자동재배분" in pg.locator(".compare > .stack").nth(0).inner_text(), "| 개정내용", pg.input_value("#summary"))
    confirm(pg); name, data = saved(pg, "t5.docx"); d = Document(io.BytesIO(data))
    exp = Document(R + "mock/예상결과_추가_Word_7장자동재배분.docx"); src = Document(R + "mock/예상결과_펀드변경_매뉴얼_v2.docx")
    print("   ", name, "| 제목 구조 목표와 같음", heads(d) == heads(exp), "| 기존 장 그대로", heads(d)[:len(heads(src))] == heads(src), "| 그림", len(src.inline_shapes), "→", len(d.inline_shapes))
    print("    개정이력", [[c.text[:16] for c in r.cells][:3] for r in d.tables[0].rows[1:]], "| 표지", [q.text for q in d.paragraphs if "버전" in q.text])
    restart(pg)

    # T6 엑셀 업무 추가
    preset(pg, "추가: 엑셀 v1 + 자동재배분 시트"); generate(pg)
    confirm(pg); name, data = saved(pg, "t6.xlsx")
    o = openpyxl.load_workbook(io.BytesIO(data)); e = openpyxl.load_workbook(R + "mock/예상결과_추가_엑셀_시트추가.xlsx")
    print("[T6 엑셀 추가]", name, "| 시트", o.sheetnames, "| 목표", e.sheetnames)
    for n in ["펀드변경", "주소변경"]: print(f"    '{n}' 원본과 같음:", snap(o[n]) == snap(s[n]))
    print("    목차 마지막 줄", [o["목차"].cell(row=7, column=c).value for c in (2, 3, 4)], "| 목표", [e["목차"].cell(row=7, column=c).value for c in (2, 3, 4)])
    restart(pg)

    # T9 개인정보
    pg.fill("#free-text", "고객 900101-1234567 요청"); pg.wait_for_timeout(300)
    print("[T9 개인정보] 경고", pg.locator("#pii").count() > 0, "| 버튼", "비활성" if pg.locator("#generate").is_disabled() else "활성", "|", pg.locator("#pii li").first.inner_text())
    pg.fill("#free-text", "")

    # T8 Q&A
    pg.click("[data-page=qa]"); pg.wait_for_selector("#question")
    print("[T8 Q&A] 검색 대상:", pg.locator("details summary.hint").inner_text())
    rows = list(csv.DictReader(open(R + "qa_test.csv", encoding="utf-8-sig")))
    for r in rows:
        pg.fill("#question", r["질문"]); pg.press("#question", "Enter"); pg.wait_for_selector("#qa-answer")
        ans = pg.locator("#qa-answer").inner_text().replace("\n", " / ")
        ev = pg.locator("#qa-evidence li").first.inner_text().split("\n")[0] if pg.locator("#qa-evidence li").count() else "-"
        print(f"  Q{r['번호']:>2} {ans[:60]} | {ev[:60]}")
    print("페이지 오류:", errors)
    b.close()
