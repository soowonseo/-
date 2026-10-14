import sys, base64, io, re
sys.path.insert(0, __import__("os").path.dirname(__file__))
from common import *
from pathlib import Path
from playwright.sync_api import sync_playwright
from docx import Document
import openpyxl
A = OUT / "artifact"
page_html = '<!doctype html><html><head><meta charset=utf8><meta name=viewport content="width=device-width,initial-scale=1"></head><body>' + (A / "manual-writer.html").read_text(encoding="utf-8") + "</body></html>"
stub = """window.__saved=null; window.claude={use: async n => n==='downloads' ? {save: async ({filename,data}) => { const b=new Uint8Array(await data.arrayBuffer()); let s=''; for (let i=0;i<b.length;i+=0x8000) s+=String.fromCharCode.apply(null,b.subarray(i,i+0x8000)); window.__saved={filename,b64:btoa(s)}; return {status:'saved'}; }} : null};"""
EXP = {"addWord": "mock/예상결과_추가_Word_7장자동재배분.docx", "addExcel": "mock/예상결과_추가_엑셀_시트추가.xlsx", "updWord": "mock/예상결과_펀드변경_매뉴얼_v2.docx",
       "updExcel": "mock/예상결과_업무매뉴얼_v2.xlsx", "newWord": "mock/예상결과_펀드변경_매뉴얼_v2.docx", "newExcel": "mock/예상결과_펀드변경_매뉴얼_v2_단독.xlsx"}
def dtexts(doc):
    out = [("P", p.style.name if p.style is not None else "-", p.text) for p in doc.paragraphs]
    out += [("T", i, [[c.text for c in r.cells] for r in t.rows]) for i, t in enumerate(doc.tables)]
    return out
def snap(ws):
    return ({c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None},
            [(i.anchor._from.row, i.anchor._from.col) for i in ws._images])
with sync_playwright() as p:
    b = launch(p); pg = b.new_page(viewport={"width": 1300, "height": 900})
    errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    route_libs(pg)
    pg.route("http://local.test/", lambda r: r.fulfill(body=page_html, content_type="text/html")); pg.add_init_script(stub)
    pg.goto("http://local.test/"); pg.wait_for_selector("#generate")
    for ex in ["addWord", "addExcel", "updWord", "updExcel", "newWord", "newExcel"]:
        pg.click("#nav-examples"); pg.click(f"button[data-ex={ex}]"); pg.click(f"button[data-ex-fill={ex}]"); pg.wait_for_selector("#generate")
        pg.click("#generate"); pg.wait_for_selector(".compare", timeout=30000)
        ver = pg.locator("input[name=verKind]:checked").evaluate_all("e=>e.map(x=>x.parentElement.innerText)")
        summ = pg.input_value("#summary"); pg.click("#confirm"); pg.wait_for_selector("#download", timeout=60000)
        pg.evaluate("window.__saved=null"); pg.click("#download"); pg.wait_for_function("window.__saved")
        s = pg.evaluate("window.__saved"); data = base64.b64decode(s["b64"]); (A / "out" / f"ex_{ex}_{s['filename']}").write_bytes(data)
        print(f"\n[{ex}] {s['filename']} | 버전선택 {ver} | 개정내용 {summ[:40]}")
        if s["filename"].endswith(".docx"):
            a, x = dtexts(Document(io.BytesIO(data))), dtexts(Document(R + EXP[ex]))
            print("  그림", len(Document(io.BytesIO(data)).inline_shapes), "/", len(Document(R + EXP[ex]).inline_shapes), "| 항목 수", len(a), "/", len(x))
            diffs = [(i, u, v) for i, (u, v) in enumerate(zip(a, x)) if u != v]
            for d in diffs[:8]: print("  다름:", str(d[1])[:150], "\n     예시:", str(d[2])[:150])
            if not diffs and len(a) == len(x): print("  내용 완전히 같음")
        else:
            o, e = openpyxl.load_workbook(io.BytesIO(data)), openpyxl.load_workbook(R + EXP[ex])
            print("  시트", o.sheetnames, "/", e.sheetnames)
            for n in e.sheetnames:
                if n not in o.sheetnames: print("  없는 시트", n); continue
                u, v = snap(o[n]), snap(e[n])
                dc = sorted(k for k in set(u[0]) | set(v[0]) if u[0].get(k) != v[0].get(k))
                print(f"  '{n}' 다른 칸 {dc} {[ (u[0].get(k), v[0].get(k)) for k in dc][:3]} | 그림위치 같음 {u[1]==v[1]}")
        pg.click("#restart")
    print("\n오류:", errs); b.close()
