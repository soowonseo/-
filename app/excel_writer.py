"""매뉴얼 마크다운 → 엑셀 시트. 배치 규칙은 PRD 6장 7을 따른다.

같은 시트 작성 함수(write_manual_sheet)를 신규·갱신·추가 세 경우에 함께 쓴다.
지금은 신규(새 통합문서에 시트 1개)만 연결되어 있다.
"""
import io
import math
import re
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Font, PatternFill
from PIL import Image as PILImage

import manual_md

TITLE_COLOR = "1F3864"
META_COLOR = "888888"
QUERY_FILL = PatternFill("solid", fgColor="F2F2F2")
CAPTURE_FILL = PatternFill("solid", fgColor="D9D9D9")
CHECK_FILL = PatternFill("solid", fgColor="FFFF00")
BOX_COLUMNS = "BCDEFGHI"
TEXT_COLUMN = "B"
CAPTION_COLUMN = "K"
IMAGE_WIDTH_PX = 550
PX_PER_ROW = 20
ROW_HEIGHT_PT = 15
TABLE_COLUMN_STEP = 3  # 표의 열은 B, E, H … 로 띄워 쓴다
SHEET_NAME_BAD = re.compile(r"[\\/*?:\[\]]")


def sheet_name_for(task_name: str) -> str:
    """'변액보험 펀드변경' → '펀드변경'. 엑셀 시트 이름에 쓸 수 없는 문자는 뺀다(최대 31자)."""
    name = (task_name or "매뉴얼").split()[-1]
    return SHEET_NAME_BAD.sub("", name)[:31] or "매뉴얼"


class _SheetWriter:
    def __init__(self, ws, image_dir: Path):
        self.ws = ws
        self.image_dir = image_dir
        self.row = 1

    def cell(self, text, col=TEXT_COLUMN, **font):
        c = self.ws[f"{col}{self.row}"]
        c.value = text
        if font:
            c.font = Font(**font)
        if manual_md.CHECK_NEEDED in str(text):
            c.fill = CHECK_FILL
        return c

    def box(self, text, fill, font=None):
        for col in BOX_COLUMNS:
            self.ws[f"{col}{self.row}"].fill = fill
        c = self.ws[f"{TEXT_COLUMN}{self.row}"]
        c.value = text
        if font:
            c.font = font
        self.row += 1

    def image(self, name, caption):
        path = self.image_dir / name
        if not path.exists():
            self.box(f"[캡처 필요] {caption or name}", CAPTURE_FILL)
            return
        with PILImage.open(path) as im:
            w, h = im.size
        img = XlImage(str(path))
        img.width = IMAGE_WIDTH_PX
        img.height = round(h * IMAGE_WIDTH_PX / w)
        self.ws.add_image(img, f"{TEXT_COLUMN}{self.row}")
        if caption:
            self.ws[f"{CAPTION_COLUMN}{self.row}"] = f"← {caption}"
        self.row += math.ceil(img.height / PX_PER_ROW) + 1

    def table(self, rows):
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
        for i, row in enumerate(cells):
            for j, value in enumerate(row):
                col = openpyxl.utils.get_column_letter(2 + j * TABLE_COLUMN_STEP)
                self.cell(value, col=col, **({"bold": True, "size": 12} if i == 0 else {}))
            self.row += 1

    def text(self, text):
        lines = manual_md.strip_source_tags(text).splitlines()
        i = 0
        while i < len(lines):
            stripped = lines[i].strip()
            if stripped.startswith("```"):
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    self.box(lines[i].rstrip(), QUERY_FILL, Font(name="Consolas", size=10))
                    i += 1
                self.row += 1
            elif stripped.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i])
                    i += 1
                self.table(rows)
                continue
            elif (m := manual_md.TOKEN_RE.fullmatch(stripped)):
                kind, value = m.groups()
                if kind == "IMAGE":
                    name, _, caption = value.partition("|")
                    self.image(name.strip(), caption.strip())
                else:
                    self.box(f"[캡처 필요] {value.strip()}", CAPTURE_FILL)
            elif stripped:
                self.cell(stripped)
                self.row += 1
            i += 1


def write_manual_sheet(ws, *, title: str, version: int, revision: tuple, history: list[str],
                       sections: list[dict], image_dir: Path) -> None:
    """시트 하나에 매뉴얼을 쓴다. 상단 2~4행은 표지·개정이력(PRD 6장 1)."""
    ver, day, summary, author, reviewer = revision
    w = _SheetWriter(ws, image_dir)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions[CAPTION_COLUMN].width = 40
    ws["B2"] = title
    ws["B2"].font = Font(bold=True, size=16, color=TITLE_COLOR)
    meta = [f"버전 {version}", day, f"작성 {author}" if author else "", f"검토 {reviewer}" if reviewer else "", "※교육용 가상 자료"]
    ws["B3"] = " / ".join(x for x in meta if x)
    ws["B3"].font = Font(size=9, color=META_COLOR)
    ws["B4"] = "개정이력: " + " → ".join(history + [f"v{ver} {day} {summary}"])
    ws["B4"].font = Font(size=9, color=META_COLOR)

    w.row = 6
    for s in sections:
        if s["level"] == 1:
            w.row += 1
            w.cell(f"{s['no']}. {s['title']}", bold=True, size=13, color=TITLE_COLOR)
        else:
            w.cell(f"{s['no']}. {s['title']}", bold=True, size=12)
        w.row += 1
        w.text(s["text"])

    # 행 높이를 모두 고정해야 그림과 옆 설명이 어긋나지 않는다(PRD 6장 7 주의 2).
    for r in range(1, max(ws.max_row, w.row) + 1):
        ws.row_dimensions[r].height = ROW_HEIGHT_PT


def build_new_xlsx(*, task_name: str, title: str, version: int, revision: tuple,
                   sections: list[dict], image_dir: Path) -> bytes:
    """신규 + 엑셀: 새 통합문서에 업무명 시트 1개."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name_for(task_name)
    write_manual_sheet(ws, title=title, version=version, revision=revision, history=[],
                       sections=sections, image_dir=image_dir)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
