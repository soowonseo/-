"""매뉴얼 마크다운 → Word(.docx). 회사 양식(word/가상_매뉴얼양식.docx)을 바탕으로 쓴다(PRD 6장 1·8)."""
import io
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import manual_md
from config import ROOT

TEMPLATE = ROOT / "word" / "가상_매뉴얼양식.docx"
HEADER_FILL = "DDE5F0"
QUERY_FILL = "F2F2F2"
CAPTURE_FILL = "D9D9D9"
CHECK_COLOR = RGBColor(0xC5, 0x5A, 0x11)
CAPTION_COLOR = RGBColor(0x59, 0x59, 0x59)
IMAGE_WIDTH = Cm(14)

INLINE_RE = re.compile(r"(\[확인 필요\]|\*\*[^*]+\*\*)")


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _add_runs(paragraph, text: str, bold: bool = False) -> None:
    """[확인 필요]는 주황색, **굵게**는 굵게."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part == manual_md.CHECK_NEEDED:
            run = paragraph.add_run(part)
            run.font.color.rgb = CHECK_COLOR
            run.bold = True
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part).bold = bold or None


def _box(doc, lines: list[str], fill: str, style: str | None = None, center: bool = False):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _shade(cell, fill)
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        if style:
            p.style = style
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(p, line)
    doc.add_paragraph()


def _capture_needed(doc, desc: str) -> None:
    _box(doc, [f"[캡처 필요] {desc}".strip()], CAPTURE_FILL, center=True)


def _image(doc, image_dir: Path, name: str, caption: str) -> None:
    path = image_dir / name
    if not path.exists():
        _capture_needed(doc, caption or name)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=IMAGE_WIDTH)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"[그림] {caption}" if caption else "[그림]")
    run.font.size = Pt(9)
    run.font.color.rgb = CAPTION_COLOR


def _table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    table = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            _add_runs(cell.paragraphs[0], value, bold=(i == 0))
            if i == 0:
                _shade(cell, HEADER_FILL)
    doc.add_paragraph()


def _write_text(doc, text: str, image_dir: Path) -> None:
    lines = manual_md.strip_source_tags(text).splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            _box(doc, code, QUERY_FILL, style="Code")
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows)
            continue
        elif (m := manual_md.TOKEN_RE.fullmatch(stripped)):
            kind, value = m.groups()
            if kind == "IMAGE":
                name, _, caption = value.partition("|")
                _image(doc, image_dir, name.strip(), caption.strip())
            else:
                _capture_needed(doc, value.strip())
        elif stripped.startswith("- "):
            _add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif stripped:
            _add_runs(doc.add_paragraph(), stripped)
        i += 1


def _fill_cover(doc, title: str, dept: str, version: int, revisions: list[tuple]) -> None:
    today = date.today().strftime("%Y.%m.%d")
    title_p = next(p for p in doc.paragraphs if p.style.name == "Title")
    info_p = next(p for p in doc.paragraphs if "버전" in p.text)
    for p, text in [(title_p, title), (info_p, f"{dept}  |  버전 {version}  |  {today}")]:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""

    table = doc.tables[0]
    template_row = table.rows[1]
    for values in revisions:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    table._tbl.remove(template_row._tr)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def _clear_after_revision(doc) -> None:
    """양식에서 개정이력 표 뒤의 안내 문구·빈 섹션 제목을 지운다(본문은 새로 쓴다)."""
    body = doc.element.body
    revision_tbl = doc.tables[0]._tbl
    after = False
    for el in list(body.iterchildren()):
        if el is revision_tbl:
            after = True
            continue
        if after and el.tag != qn("w:sectPr"):
            body.remove(el)


def build_docx(*, title: str, dept: str, version: int, sections: list[dict],
               image_dir: Path, revisions: list[tuple]) -> bytes:
    doc = Document(TEMPLATE)
    _fill_cover(doc, title, dept, version, revisions)
    _clear_after_revision(doc)
    for s in sections:
        doc.add_paragraph(f"{s['no']}. {s['title']}", style="Heading 1" if s["level"] == 1 else "Heading 2")
        _write_text(doc, s["text"], image_dir)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
