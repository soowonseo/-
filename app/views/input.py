import io
import re

import openpyxl
import streamlit as st
from docx import Document

import db
import generator
import storage

JOB_NEW, JOB_UPDATE, JOB_ADD = "신규 작성", "기존 매뉴얼 갱신", "기존 매뉴얼에 업무 추가"
IMAGE_TYPES = ["전산화면", "엑셀", "쿼리결과", "업무흐름도"]


def guess_image_type(file_name: str) -> str:
    """파일명으로 종류를 미리 골라 둔다. 사용자가 바꿀 수 있다."""
    for keyword, kind in [("흐름도", "업무흐름도"), ("쿼리", "쿼리결과"), ("엑셀", "엑셀")]:
        if keyword in file_name:
            return kind
    return "전산화면"


def sheet_names(file) -> list[str]:
    return openpyxl.load_workbook(io.BytesIO(file.getvalue()), read_only=True).sheetnames


def word_chapter_titles(file) -> list[str]:
    """'7. 자동재배분'처럼 번호로 시작하는 문단을 장 제목 후보로 본다."""
    doc = Document(io.BytesIO(file.getvalue()))
    return [p.text.strip() for p in doc.paragraphs if re.match(r"^\d+\.\s*\S", p.text.strip())]


st.title("📥 자료 입력")

job = st.radio("어떤 작업을 할까요?", [JOB_NEW, JOB_UPDATE, JOB_ADD], horizontal=True)

output_format = None
existing = None
source_sheet = None
task_name = None
problems = []

if job == JOB_NEW:
    choice = st.radio("어떤 형식으로 만들까요?", ["Word", "엑셀"], index=None, horizontal=True)
    if choice is None:
        problems.append("출력 형식(Word/엑셀)을 골라 주세요.")
    else:
        output_format = "docx" if choice == "Word" else "xlsx"
else:
    existing = st.file_uploader("기존 매뉴얼 (.docx 또는 .xlsx)", type=["docx", "xlsx"])
    if existing is None:
        problems.append("기존 매뉴얼을 올려 주세요.")
    else:
        output_format = existing.name.rsplit(".", 1)[-1].lower()
        st.caption(f"출력 형식: 기존과 같은 **{'Word' if output_format == 'docx' else '엑셀'}**")
        if job == JOB_UPDATE and output_format == "xlsx":
            source_sheet = st.selectbox("갱신할 시트 1개를 고르세요", sheet_names(existing))
        if job == JOB_ADD:
            task_name = st.text_input("추가할 업무명", placeholder="예: 자동재배분").strip()
            if not task_name:
                problems.append("추가할 업무명을 입력해 주세요.")
            else:
                if output_format == "xlsx":
                    taken = task_name in sheet_names(existing)
                else:
                    taken = any(task_name in t for t in word_chapter_titles(existing))
                if taken:
                    st.warning(f"'{task_name}' 이름의 {'시트가' if output_format == 'xlsx' else '장이'} 이미 있습니다.")

st.divider()
st.subheader("업무 내용")
text = st.text_area("줄글 입력 — 업무 내용을 직접 쓰거나 붙여넣으세요. 러프하게 적어도 됩니다.", height=200)

st.subheader("원자료 업로드")
col1, col2 = st.columns(2)
memos = col1.file_uploader("메모 (.txt)", type=["txt"], accept_multiple_files=True)
sheets = col1.file_uploader("엑셀 (.xlsx)", type=["xlsx"], accept_multiple_files=True)
queries = col2.file_uploader("쿼리 (.sql)", type=["sql"], accept_multiple_files=True)
images = col2.file_uploader("이미지 (.png/.jpg, 여러 장)", type=["png", "jpg", "jpeg"],
                            accept_multiple_files=True)

image_meta = []
if images:
    st.subheader("이미지 정보")
    for i, img in enumerate(images):
        with st.container(border=True):
            c_img, c_meta = st.columns([1, 3])
            c_img.image(img, width="stretch")
            kind = c_meta.selectbox("종류", IMAGE_TYPES, index=IMAGE_TYPES.index(guess_image_type(img.name)),
                                   key=f"kind_{i}_{img.name}")
            step = c_meta.text_input("배치할 단계 (선택, 예: 4-3)", key=f"step_{i}_{img.name}").strip()
            caption = c_meta.text_input("한 줄 설명", key=f"cap_{i}_{img.name}").strip()
            c_meta.caption(img.name)
            image_meta.append((img, kind, step or None, caption or None))

raw_files = [*memos, *sheets, *queries, *images]
has_raw = bool(text.strip()) or bool(raw_files)
if job == JOB_NEW and not has_raw:
    problems.append("줄글을 쓰거나 원자료 파일을 하나 이상 올려 주세요.")
if job == JOB_ADD and not has_raw:
    problems.append("추가할 업무의 줄글이나 원자료를 넣어 주세요.")

st.divider()
for p in problems:
    st.caption(f"⚠️ {p}")

if st.button("초안 생성", type="primary", disabled=bool(problems)):
    manual_id = db.create_manual(
        job_type={JOB_NEW: "신규", JOB_UPDATE: "갱신", JOB_ADD: "추가"}[job],
        task_name=task_name,
        output_format=output_format,
        source_format=output_format if existing else "없음",
        source_sheet=source_sheet,
        source_files=[f.name for f in raw_files],
    )
    if existing:
        db.set_source_file(manual_id, storage.save_upload(manual_id, existing))
    if text.strip():
        db.add_asset(manual_id, storage.save_text(manual_id, "줄글.txt", text), "줄글")
    for files, kind in [(memos, "메모"), (sheets, "엑셀"), (queries, "쿼리")]:
        for f in files:
            db.add_asset(manual_id, storage.save_upload(manual_id, f), kind)
    for img, kind, step, caption in image_meta:
        db.add_asset(manual_id, storage.save_upload(manual_id, img), kind, step, caption)

    with st.spinner("초안을 만드는 중..."):
        generator.generate(manual_id)
    st.session_state["manual_id"] = manual_id
    st.switch_page("views/review.py")
